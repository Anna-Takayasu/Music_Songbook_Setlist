"""Simple print-friendly setlist Word export. Each set gets its own page."""
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_setlist_word(songs, sets, concert_title='Setlist', show_duration=False):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(18)
    normal.paragraph_format.space_after = Pt(8)
    for n, plan in enumerate(sets):
        if n:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(concert_title or 'Setlist')
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(22)
        p.paragraph_format.space_after = Pt(20)
        heading = doc.add_paragraph()
        run = heading.add_run(plan.get('name') or f'Set {n+1}')
        run.bold = True
        run.font.size = Pt(20)
        run.font.name = 'Arial'
        heading.paragraph_format.space_after = Pt(12)
        for pos, idx in enumerate(plan.get('songs', []), 1):
            if not isinstance(idx, int) or idx < 0 or idx >= len(songs):
                continue
            song = songs[idx]
            line = f"{pos}. {song.get('title', 'Untitled')}"
            if show_duration and song.get('duration_seconds'):
                t = int(song['duration_seconds'])
                line += f'  ({t//60}:{t%60:02d})'
            p = doc.add_paragraph(line)
            p.paragraph_format.keep_together = True
        if not plan.get('songs'):
            doc.add_paragraph('(No songs selected)')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
