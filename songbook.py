"""Songbook data, DOCX importer and Word export."""
import re
import html
from capo_display import concert_chord, capo_fret
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PRESETS = []  # Public version intentionally ships no real arrangements.


# Safe lightweight lyric markup supported by both preview and Word export:
# **bold**, [color=red]...[/color], [color=#RRGGBB]...[/color].
COLOR_NAMES={'red':'C62828','blue':'1565C0','green':'2E7D32','orange':'EF6C00',
             'purple':'7B1FA2','black':'222222','gray':'666666','grey':'666666'}
def parse_lyric_runs(line):
    tokens=re.split(r'(\*\*|\[color=(?:#[0-9a-fA-F]{6}|red|blue|green|orange|purple|black|gr[ae]y)\]|\[/color\])', line,flags=re.I)
    bold=False; colors=[]; out=[]
    for token in tokens:
        if not token: continue
        if token=='**': bold=not bold
        elif token.lower().startswith('[color='):
            raw=token[7:-1].lower(); colors.append(COLOR_NAMES.get(raw,raw.lstrip('#').upper()))
        elif token.lower()=='[/color]':
            if colors:colors.pop()
        else:out.append((token,bold,colors[-1] if colors else None))
    return out

def add_formatted_lyric(p,line,size=9):
    for content,bold,color in parse_lyric_runs(line):
        run=p.add_run(content);run.bold=bold;run.font.size=Pt(size)
        if color:run.font.color.rgb=RGBColor.from_string(color)

def lyric_html(lyrics):
    pieces=[]
    for line in lyrics.splitlines():
        if not line.strip():pieces.append('<div style="height:10px"></div>');continue
        current=''
        for txt,bold,color in parse_lyric_runs(line):
            style=('font-weight:700;' if bold else '')+(f'color:#{color};' if color else '')
            current+=f'<span style="{style}">{html.escape(txt)}</span>'
        pieces.append(f'<div>{current}</div>')
    return ''.join(pieces)

def preview_html(song):
    header=html.escape(song.get('title','Untitled'))
    key=html.escape(song.get('key',''))
    parts=[f'<section style="font-family:Arial,sans-serif;background:white;color:#222;padding:22px;max-width:780px;border:1px solid #ddd;border-radius:8px"><h2 style="margin:0 0 4px">{header} <small style="font-size:55%;color:#536678">({key})</small></h2>']
    parts.append(f'<div style="color:#555;font-size:12px;margin-bottom:10px">{html.escape(song.get("artist", ""))} · {html.escape(song.get("kind", "Song"))} · Capo: {html.escape(str(song.get("capo", "None")))}</div>')
    if capo_fret(song):parts.append('<div style="font-size:11px;color:#516575;margin-bottom:6px">Bold = guitar shape · blue parentheses = sounding chord (without capo)</div>')
    if song.get('structure'):parts.append(f'<div style="font-size:11px;margin-bottom:12px"><b>STRUCTURE</b> {html.escape(song["structure"])}</div>')
    for sec in song.get('sections',[]):
        parts.append(f'<h3 style="font-size:14px;margin:13px 0 5px">{html.escape(sec.get("name", ""))}</h3>')
        for line in rows(sec.get('chords','')):
            notes=[c for c in line if c!='[2/4]']
            parts.append('<div style="display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:3px;margin:3px 0">')
            for ch in notes:
                actual=concert_chord(ch,song)
                suffix=f' <span style="color:#1864ab;font-size:85%;font-weight:normal">({html.escape(actual)})</span>' if actual else ''
                parts.append(f'<div style="border:1px solid #bdcbd9;background:#edf2f9;text-align:center;padding:6px;font-weight:bold;border-radius:3px">{html.escape(ch)}{suffix}</div>')
            parts.append('</div>')
            if '[2/4]' in line:parts.append('<small>2/4 short bar</small>')
    lyric= song.get('lyrics','').strip()
    if lyric:
        mode=song.get('lyrics_layout','Auto');count=sum(_lyric_line_count(x) for x in _lyrics_blocks(lyric))
        two=mode=='Two columns' or (mode!='One column' and (count>13 or len(lyric)>580))
        parts.append('<h3 style="font-size:14px">Lyrics</h3>')
        parts.append(f'<div style="font-size:12px;line-height:1.45;white-space:normal;column-count:{2 if two else 1};column-gap:24px">{lyric_html(lyric)}</div>')
    parts.append('</section>');return ''.join(parts)

def rows(chords):
    """Space or | separated tokens. One chord token occupies one 2-beat box.
    Use '-' as a hold, and [2/4] at end of row for a shortened bar annotation."""
    return [[x.strip() for x in re.split(r'\s+|\s*\|\s*',line.strip()) if x.strip()] for line in chords.splitlines() if line.strip()]

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)

def _lyrics_blocks(lyrics):
    # Respect stanza separators where present. A paragraph per line from some importers
    # is treated as a line so it can be balanced between columns.
    cleaned=lyrics.strip().replace('\r\n','\n')
    return [p.strip() for p in re.split(r'\n\s*\n', cleaned) if p.strip()]

def _lyric_line_count(text):
    return sum(max(1, (len(line)+48)//49) for line in text.splitlines())

def _write_lyrics(doc, song, compact=False):
    lyric= song.get('lyrics','').strip()
    if not lyric:return
    doc.add_heading('Lyrics',level=2)
    blocks=_lyrics_blocks(lyric)
    total_lines=sum(_lyric_line_count(b) for b in blocks)
    mode=song.get('lyrics_layout','Auto')
    use_two=mode=='Two columns' or (mode!='One column' and (total_lines>13 or len(lyric)>580))
    if use_two:
        table=doc.add_table(rows=1,cols=2)
        table.autofit=False
        table.columns[0].width=Inches(3.32)
        table.columns[1].width=Inches(3.32)
        split_target=sum(_lyric_line_count(b) for b in blocks)/2
        left=[];right=[];length=0
        for block in blocks:
            if length<split_target or not left:
                left.append(block);length+=_lyric_line_count(block)
            else:right.append(block)
        # In case the input consists of a single long stanza, divide on lines.
        if not right and left:
            all_lines=left[0].splitlines()
            if len(all_lines)>1:
                mid=(len(all_lines)+1)//2
                left=['\n'.join(all_lines[:mid])]
                right=['\n'.join(all_lines[mid:])]
        for cell, items in zip(table.rows[0].cells,[left,right]):
            cell.width=Inches(3.32)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text=''
            for idx,blk in enumerate(items):
                p=cell.paragraphs[0] if idx==0 else cell.add_paragraph()
                p.paragraph_format.space_after=Pt(6 if compact else 8)
                for j,line in enumerate(blk.splitlines()):
                    if j:p.add_run().add_break()
                    add_formatted_lyric(p,line,9 if compact else 9.5)
    else:
        for verse in blocks:
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(5 if compact else 8)
            for ix,line in enumerate(verse.splitlines()):
                if ix:p.add_run().add_break()
                add_formatted_lyric(p,line,9 if compact else 9.5)

def export_docx(songs, heading='BLUEGRASS SONGBOOK'):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62)
    sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    normal=doc.styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(4)
    for st_name,size in [('Heading 1',16),('Heading 2',11)]:
        st=doc.styles[st_name];st.font.size=Pt(size)
        st.paragraph_format.space_before=Pt(6);st.paragraph_format.space_after=Pt(3)
    title=doc.add_heading(heading,0);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('SONGS  |  INSTRUMENTALS');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('Contents',level=1)
    ordered=sorted(songs,key=lambda s:(s.get('kind','Song')=='Instrumental',s['title'].casefold()))
    for kind in ['Song','Instrumental']:
        subset=[s for s in ordered if s.get('kind','Song')==kind]
        doc.add_heading('Songs' if kind=='Song' else 'Instrumentals',level=2)
        for s in subset:doc.add_paragraph(f"{s['title']}  ({s.get('key','')})",style='List Bullet')
    for s in ordered:
        doc.add_page_break()
        compact=bool(s.get('lyrics','').strip())
        doc.add_heading(f"{s['title']}  ({s.get('key','')})",level=1)
        p=doc.add_paragraph();p.add_run(s.get('artist','')).italic=True
        doc.add_paragraph('  |  '.join(x for x in [s.get('kind','Song'),f"Capo: {s.get('capo','None')}",f"Tempo: {s.get('bpm','')} BPM" if s.get('bpm') else '2 beats / box'] if x))
        if capo_fret(s):
            lg=doc.add_paragraph('Bold = guitar shape  |  blue parentheses = sounding chord (without capo)')
            for run in lg.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(78,96,113)
        if s.get('structure'):doc.add_paragraph('STRUCTURE  '+s['structure'])
        if s.get('notes'):doc.add_paragraph('NOTE  '+s['notes'])
        for part in s.get('sections',[]):
            doc.add_heading(part.get('name','Section'),level=2)
            for line in rows(part.get('chords','')):
                short='[2/4]' in line;line=[ch for ch in line if ch!='[2/4]']
                for i in range(0,len(line),4):
                    chunk=line[i:i+4];table=doc.add_table(rows=1,cols=4)
                    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
                    for j,cell in enumerate(table.rows[0].cells):
                        cell.width=Inches(1.6);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        guitar=chunk[j] if j<len(chunk) else ''
                        cell.text=''
                        p=cell.paragraphs[0]
                        if guitar:
                            lead=p.add_run(guitar);lead.bold=True;lead.font.size=Pt(10.5 if compact else 12)
                            actual=concert_chord(guitar,s)
                            if actual:
                                sub=p.add_run(f' ({actual})');sub.bold=False;sub.font.size=Pt(8 if compact else 9)
                                sub.font.color.rgb=RGBColor(24,100,171)
                        shade(cell,'EAF0F7' if j<2 else 'F3F6FA')
                        for para in cell.paragraphs:
                            para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            para.paragraph_format.space_after=Pt(0)
                            # Individual guitar and concert-pitch runs have different styling.
                    # No empty spacer paragraph after each row; keeps chords compact.
                if short:doc.add_paragraph('Short measure (2/4)',style='Caption')
        _write_lyrics(doc,s,compact=compact)
    return doc

def import_docx(path):
    """Best-effort import for existing chord-table Word songbooks.
    Gets headings, tables and lyrics in their original order; original beat values may be unknown."""
    from docx.oxml.ns import qn
    doc=Document(path); out=[]; current=None; current_part=None
    for child in doc.element.body.iterchildren():
        if child.tag==qn('w:p'):
            from docx.text.paragraph import Paragraph
            p=Paragraph(child,doc); t=p.text.strip()
            if not t: continue
            if p.style.name.startswith('Heading 1'):
                if current: out.append(current)
                m=re.match(r'^(.*?)\s*\(([^()]*)\)\s*$',t)
                title,key=(m.group(1),m.group(2)) if m else (t,'')
                current={"title":title,"artist":"","key":key,"kind":"Song","capo":"","bpm":"","notes":"Imported from original DOCX. Check beat durations, section alignment, and alternate-key tables.","structure":"","lyrics":"","sections":[]};current_part=None
            elif current:
                if len(t)<38 and ('verse' in t.lower() or 'chorus' in t.lower() or 'intro' in t.lower() or 'bridge' in t.lower() or 'part' in t.lower()):
                    current_part={"name":t,"chords":""};current['sections'].append(current_part)
                elif t.strip('- ').strip():
                    current['lyrics'] += ('\n\n' if current['lyrics'] else '')+t
        elif child.tag==qn('w:tbl') and current:
            from docx.table import Table
            tab=Table(child,doc)
            lines=[' '.join(c.text.strip().replace('\n','/') or '-' for c in row.cells) for row in tab.rows]
            if current_part is None:
                current_part={"name":"Chords","chords":""};current['sections'].append(current_part)
            current_part['chords']+='\n'.join(lines)+'\n'
    if current:out.append(current)
    instrumentals={'Cherokee Shuffle','Blackberry Blossom','Birdland breakdown','Tennessee politics','High road','Salt Spring','Clinch Mountain Backstep'}
    for s in out:
        if s['title'].casefold() in {x.casefold() for x in instrumentals}:s['kind']='Instrumental'
    return out

def synth_preview(song, max_seconds=150):
    """Basic synthetic plucked-chord rehearsal WAV (no melody), stdlib only."""
    import math, wave, io, struct
    notes={'C':0,'C#':1,'Db':1,'D':2,'D#':3,'Eb':3,'E':4,'F':5,'F#':6,'Gb':6,'G':7,'G#':8,'Ab':8,'A':9,'A#':10,'Bb':10,'B':11}
    sr=12000; bpm=float(song.get('bpm') or 105); duration=120/bpm
    samples=[]; previous=None
    for section in song.get('sections',[]):
      for line in rows(section.get('chords','')):
        for ch in line:
          if ch=='[2/4]':continue
          if ch!='-':previous=ch
          else:ch=previous or 'G'
          m=re.match(r'^([A-G](?:#|b)?)(.*?)(?:/[A-G][#b]?)?$',ch)
          if not m:continue
          root,quality=m.groups()
          intervals=[0,3,7] if quality.startswith('m') and not quality.startswith('maj') else [0,4,7]
          if '7' in quality:intervals.append(10 if 'maj7' not in quality else 11)
          base=48+notes.get(root,0)
          count=min(int(sr*duration),int(max_seconds*sr)-len(samples))
          if count<=0:break
          freqs=[440*2**((base+n-69)/12) for n in intervals]
          for k in range(count):
            t=k/sr; beat=t*bpm/60
            envelope=math.exp(-2.5*(t%duration)) * min(1, t*100)
            # rhythmic chord pulses on beats 1 and 2 per box
            p=max(math.exp(-18*(t-0.02)) if t>=.02 else 0,math.exp(-18*(t-60/bpm)) if t>=60/bpm else 0)
            value=sum(math.sin(2*math.pi*f*t)+.18*math.sin(2*math.pi*2*f*t) for f in freqs)/len(freqs)
            samples.append(int(15000*value*(.28*envelope+.35*p)))
        if len(samples)>=max_seconds*sr:break
    buf=io.BytesIO()
    with wave.open(buf,'wb') as wav:
      wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(sr)
      wav.writeframes(struct.pack('<'+'h'*len(samples),*samples))
    return buf.getvalue()
